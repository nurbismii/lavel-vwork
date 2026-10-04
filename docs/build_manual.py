from pathlib import Path
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / 'Manual_Pengguna_dan_Rumus_Beban_Kerja_RuangKerja.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.1)
sec.footer_distance = Cm(1)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('RuangKerja  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)
doc.core_properties.title = 'Manual Pengguna dan Rumus Perhitungan Beban Kerja RuangKerja'
doc.core_properties.subject = 'Panduan operasional pencatatan aktivitas aktual dan analisis utilisasi'
doc.core_properties.author = 'Dokumentasi RuangKerja'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def steps(items):
    for i, text in enumerate(items, 1): p(f'{i}. {text}')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for c, val in zip(t.rows[0].cells, headers): c.text = val
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, val in zip(cells, row): c.text = str(val)
    for ri, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            tc = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DBE5F1' if ri == 0 else ('F5F7FA' if ri % 2 == 0 else 'FFFFFF')); tc.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '90'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tc.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            tc.append(borders)
            valign = OxmlElement('w:vAlign'); valign.set(qn('w:val'), 'center'); tc.append(valign)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    p('') .paragraph_format.space_after = Pt(0)

p('Manual Pengguna RuangKerja', 'Title')
p('Panduan operasional dan rumus perhitungan beban kerja', 'Subtitle')
p('Versi dokumen 1.0  |  30 September 2026')
p('Manual ini membantu Staff, Manager, PIC HR/Operasional, Administrator, dan Manajemen/Viewer menggunakan RuangKerja: menyiapkan periode, mencatat pekerjaan, memvalidasi laporan, membaca utilisasi, dan menindaklanjuti hasilnya.')
p('Beban kerja pada alur input saat ini dihitung dari total durasi aktivitas aktual dibandingkan dengan kapasitas efektif. Persentase progres pekerjaan merupakan informasi penyelesaian dan tidak menjadi pengali durasi beban kerja.')
h('Panduan membaca')
table(['Bagian', 'Isi'], [
('1', 'Akses aplikasi dan pembagian peran'),
('2', 'Persiapan organisasi dan periode'),
('3', 'Pencatatan aktivitas oleh Staff'),
('4', 'Progres pekerjaan dan pengajuan'),
('5', 'Validasi dan tindak lanjut'),
('6', 'Dashboard dan laporan'),
('7', 'Rumus hari kerja dan kapasitas efektif'),
('8', 'Rumus beban aktual dan klasifikasi'),
('9', 'Contoh hitung individu dan tim'),
('10', 'Pemeriksaan hasil dan penanganan kendala'),
('Lampiran', 'Rujukan implementasi dan istilah data'),
], [2, 14.8])
p('Angka contoh dalam manual adalah ilustrasi, bukan data pegawai. Ambang 70%, 85%, dan 100% merupakan nilai cadangan aplikasi; kebijakan organisasi yang berlaku dapat berbeda.')

page('1 Akses aplikasi dan pembagian peran')
h('Masuk dan menjaga akses')
steps(['Buka alamat RuangKerja yang diberikan Administrator, kemudian halaman login.',
'Masukkan email dan kata sandi akun aktif. Akun dibuat Administrator; pendaftaran publik tidak tersedia.',
'Setelah masuk, periksa nama akun dan menu yang tersedia. Pilih periode yang ingin dilihat pada Dashboard.',
'Gunakan menu Profil melalui avatar untuk memperbarui profil atau mengganti kata sandi. Keluar setelah menggunakan perangkat bersama.'])
h('Hak akses utama')
table(['Peran', 'Cakupan data aktif', 'Tugas utama'], [
('Staff', 'Diri sendiri', 'Input aktivitas, progres, dan pengajuan'),
('Manager', 'Bawahan langsung menurut relasi atasan', 'Pantau tim, validasi, kelola tindak lanjut'),
('PIC HR/Operasional', 'Pengguna dalam unitnya', 'Validasi unit, periode, ambang, master, audit'),
('Administrator', 'Seluruh pengguna aktif', 'Organisasi, akun, konfigurasi, dan validasi'),
('Manajemen/Viewer', 'Pengguna dalam unitnya', 'Pantau dashboard dan laporan'),
], [3.5, 5.3, 8])
p('Semua pengguna yang mendapat penugasan dapat membuka Tindak lanjut saya. Hak mengelola penugasan tersedia untuk Manager, PIC HR/Operasional, dan Administrator. Menu Input saya hanya tersedia untuk Staff.')
h('Navigasi yang digunakan')
p('Dashboard menampilkan ringkasan dan anggota. Input saya membuka Laporan kerja saya. Validasi dipakai untuk keputusan pengajuan. Kelola tindak lanjut dipakai untuk penugasan. Periode, Ambang utilisasi, Master aktivitas, Audit, dan Organisasi muncul sesuai hak akses.')
p('Relasi atasan menentukan anggota yang terlihat oleh Manager; posisi atau jabatan saja tidak memberikan hak validasi. Untuk alur atasan gunakan peran Manager. Peran Supervisor yang tercantum dalam kode belum memiliki pemetaan lengkap pada layanan cakupan data dan tidak digunakan dalam prosedur manual ini.')

page('2 Persiapan organisasi dan periode')
h('Penyiapan oleh Administrator')
steps(['Buka Organisasi. Buat unit terlebih dahulu, lalu buat pengguna dan lengkapi identitas, jabatan, email, peran, unit, atasan, serta status aktif.',
'Isi hari kerja dalam siklus, hari off, jam kerja per hari, dan hari pertama siklus kerja. Tanggal acuan dihitung sebagai hari kerja pertama, bukan hari off.',
'Periksa relasi Staff ke Manager. Pastikan jadwal benar sebelum Staff membuka input untuk pertama kali pada periode tersebut.'])
h('Membuat dan membuka periode')
steps(['PIC atau Administrator membuka Periode dan membuat periode bulanan dengan tanggal pertama bulan, tenggat pengajuan, persentase produktif, dan catatan kebijakan bila diperlukan.',
'Saat status masih Draft, periksa dan koreksi data periode serta standar persentase produktif. Persentase harus lebih dari 0% sampai dengan 100%.',
'Buka periode. Hanya satu periode dapat berstatus terbuka; selesaikan dan kunci periode aktif sebelum membuka periode berikutnya.',
'Setelah pengisian dan validasi selesai, kunci periode. Bila koreksi diperlukan, gunakan Buka kembali dan isi alasan minimal 10 karakter. Pastikan tidak ada periode lain yang terbuka.'])
p('Pengajuan pertama baru tersedia setelah akhir bulan. Karena itu, atur tenggat agar memberi waktu pengajuan dan validasi setelah bulan selesai. Tenggat tidak otomatis mengunci periode; penguncian merupakan tindakan pengelola.')
h('Ambang utilisasi dan master aktivitas')
p('Pada Ambang utilisasi, isi batas Kapasitas tersedia, Sehat, dan Padat secara berurutan naik, tanggal berlaku, penanda sementara, serta alasan perubahan. Tanggal berlaku tidak boleh sebelum hari ini. Ambang yang dipilih sistem mengikuti tanggal awal periode, bukan tanggal saat laporan dibuka.')
p('Pada Master aktivitas, rapikan kategori dan sifat pekerjaan; nonaktifkan opsi yang tidak dipakai. Penggabungan opsi duplikat mengalihkan pemakaian aktivitas lama ke tujuan. Tinjau dampaknya terhadap laporan historis sebelum menggabungkan.')
p('Kapasitas disimpan sebagai salinan kondisi jadwal saat pertama dibuat. Mengubah jadwal profil tidak otomatis menghitung ulang kapasitas yang sudah tersimpan. Koordinasikan koreksi historis dengan pengelola aplikasi.')

page('3 Pencatatan aktivitas oleh Staff')
steps(['Buka Input saya. Pastikan periode dan Acuan kapasitas sesuai. Jika belum ada periode terbuka, minta PIC membuka periode.',
'Pada Catat aktivitas aktual, isi tanggal aktivitas, nama, kategori, sifat pekerjaan, dan durasi aktual dalam menit.',
'Isi hasil/volume dan satuan bila diperlukan. Tambahkan keterangan hasil atau kendala untuk membantu validasi.',
'Tekan Catat aktivitas aktual. Pastikan pesan berhasil muncul dan catatan masuk ke daftar. Periksa perubahan total waktu serta utilisasi.'])
table(['Isian', 'Ketentuan'], [
('Tanggal aktivitas', 'Dalam bulan yang terbuka dan tidak melewati hari ini.'),
('Nama aktivitas', 'Wajib, maksimal 255 karakter; gunakan nama yang jelas dan konsisten.'),
('Kategori dan sifat pekerjaan', 'Wajib, masing-masing maksimal 100 karakter. Pilih master atau ketik opsi baru.'),
('Durasi aktual', 'Wajib, bilangan bulat 1–1.440 menit untuk satu catatan.'),
('Hasil atau volume', 'Opsional. Jika diisi harus lebih besar dari nol; satuan menjadi wajib.'),
('Satuan dan keterangan', 'Satuan maksimal 50 karakter; keterangan maksimal 1.000 karakter.'),
], [4.4, 12.4])
h('Contoh pencatatan yang benar')
p('Meninjau 3 dokumen membutuhkan total 45 menit. Masukkan durasi 45, volume 3, dan satuan Dokumen. Beban bertambah 45 menit. Jangan memasukkan 135 menit karena durasi sudah merupakan total waktu untuk ketiga dokumen.')
p('Jika volume dikosongkan, sistem menggunakan volume 1; bila satuan juga kosong, satuannya menjadi Aktivitas. Volume tidak wajib untuk menghitung utilisasi.')
h('Koreksi dan ketepatan waktu')
p('Form saat ini menyediakan tambah dan hapus aktivitas. Untuk salah input, selama status Draft atau Perlu Revisi dan periode terbuka, hapus catatan yang keliru lalu catat ulang dengan benar. Periksa kembali pelengkap progres yang terkait sebelum mengajukan.')
p('Tanggal penginputan ditentukan server. Pencatatan pada hari aktivitas diberi label Tepat waktu; pencatatan pada hari berikutnya diberi label Terlambat sejumlah hari. Label keterlambatan tidak mengalikan atau mengurangi durasi.')
p('Hindari mencatat menit yang sama dua kali pada pekerjaan yang berlangsung bersamaan. Batas 1.440 menit berlaku per catatan; pengguna tetap perlu meninjau kewajaran total harian.')

page('4 Progres pekerjaan dan pengajuan')
h('Melengkapi pekerjaan aktual')
steps(['Pada Lengkapi laporan progres, pilih tanggal laporan dan pekerjaan dari aktivitas aktual yang sudah tersimpan.',
'Pilih Selesai dikerjakan atau Sedang dikerjakan. Isi ringkasan hasil, kendala bila ada, serta langkah yang sudah atau akan diambil.',
'Untuk selesai, isi progres 100%. Untuk sedang dikerjakan, isi 1–99% dan target selesai.',
'Tekan Simpan pelengkap laporan. Penyimpanan dengan tanggal, kategori, dan nama yang sama memperbarui catatan progres pada tanggal tersebut.'])
h('Mencatat rencana pekerjaan')
p('Gunakan formulir rencana untuk pekerjaan yang belum dimulai. Isi tanggal laporan, target mulai atau selesai, kategori, nama, ringkasan rencana, dan langkah berikutnya. Sistem menyimpan status Akan dikerjakan dengan progres 0%. Rencana tidak menambah menit beban kerja.')
p('Tanggal laporan harus berada dalam periode terbuka dan tidak di masa depan. Target merupakan tanggal rencana, sehingga dapat berada setelah tanggal laporan. Saat pekerjaan benar-benar dilakukan, catat aktivitas aktual dan lengkapi progresnya.')
h('Mengajukan ringkasan bulanan')
steps(['Pastikan kapasitas tersedia dan minimal satu aktivitas aktual sudah tercatat. Tinjau durasi, volume, kategori, serta keterangan.',
'Setelah bulan berakhir dan periode masih terbuka, isi Catatan untuk atasan bila diperlukan, lalu tekan Ajukan untuk validasi.',
'Status menjadi Diajukan. Input dan penghapusan terkunci selama pengajuan menunggu keputusan.',
'Jika diminta revisi, baca catatan reviewer, perbaiki data ketika periode terbuka, lalu pilih Ajukan ulang untuk validasi. Pengajuan ulang berstatus Perlu Revisi tidak menunggu akhir bulan lagi.'])
table(['Status pengajuan', 'Boleh mengubah aktivitas'], [
('Draft', 'Ya, bila periode terbuka'),
('Diajukan', 'Tidak'),
('Perlu Revisi', 'Ya, bila periode terbuka'),
('Disetujui', 'Tidak; reviewer harus meminta revisi dahulu'),
], [5, 11.8])
p('Draft formulir lokal belum berarti data tersimpan di server. Pastikan ada pesan berhasil dan catatan muncul sebelum meninggalkan halaman.')

page('5 Validasi dan tindak lanjut')
h('Validasi oleh Manager atau pengelola')
steps(['Buka Validasi dan pilih pengajuan dalam cakupan akses. Periksa periode, pemilik, kapasitas efektif, total aktual, dan utilisasi.',
'Bandingkan catatan aktivitas dengan konteks pekerjaan. Periksa duplikasi, kewajaran menit, kelengkapan pekerjaan, dan alasan beban tinggi atau rendah.',
'Pilih Disetujui bila data dapat diterima, atau Perlu Revisi bila perlu perbaikan. Catatan wajib untuk permintaan revisi; tulis bagian yang harus diperbaiki.',
'Setelah Staff mengajukan ulang, tinjau kembali lalu berikan keputusan. Pengajuan yang sudah disetujui dapat dikembalikan menjadi Perlu Revisi dengan catatan.'])
p('Alur normal adalah Draft → Diajukan → Disetujui. Alur perbaikan adalah Diajukan → Perlu Revisi → Diajukan. Perubahan keputusan tercatat bersama pelaku dan catatan validasi.')
h('Membuat tindakan perbaikan')
steps(['Buka Kelola tindak lanjut. Tentukan unit dan owner atau penanggung jawab yang sesuai cakupan akses.',
'Pilih jenis tindakan: redistribusi pekerjaan, perbaikan proses, otomasi, atau kajian tenaga kerja. Isi judul dan uraian yang menyebut masalah serta hasil yang diharapkan.',
'Tautkan pengajuan bila relevan, isi target tanggal bila diperlukan, dan simpan dengan status yang sesuai.',
'Pantau penyelesaian. Owner membuka Tindak lanjut saya untuk memperbarui status tugas miliknya. Tugas yang dibatalkan tidak dapat diperbarui oleh owner.'])
table(['Status tindak lanjut', 'Makna operasional'], [
('open', 'Terbuka atau belum dimulai'),
('in_progress', 'Sedang dilaksanakan'),
('completed', 'Selesai'),
('cancelled', 'Dibatalkan oleh pengelola'),
], [4.8, 12])
p('Contoh: utilisasi tinggi akibat laporan berulang dapat ditindaklanjuti dengan perbaikan format atau otomasi. Tentukan owner dan tanggal target agar keputusan dapat dipantau. Jangan menyimpulkan kebutuhan penambahan pegawai hanya dari satu persentase tanpa memeriksa kualitas catatan dan periode pembanding.')

page('6 Dashboard dan laporan')
h('Membaca Dashboard')
steps(['Pilih periode, lalu filter unit dan status beban jika diperlukan. Filter unit tetap dibatasi hak akses.',
'Baca kapasitas efektif, waktu aktual, utilisasi tim, dan distribusi status. Angka ringkasan tim berasal dari pengajuan Disetujui dalam hasil filter.',
'Buka detail anggota untuk menelusuri aktivitas, kapasitas, dan status pengajuan. Anggota yang belum mengisi atau belum memiliki kapasitas tidak boleh langsung disimpulkan memiliki beban nol.'])
p('Ringkasan tim memakai perbandingan jumlah waktu terhadap jumlah kapasitas. Persentase individu tidak dirata-ratakan secara sederhana. Filter status dapat mengubah kelompok anggota dan angka ringkasan yang dihitung.')
h('Ekspor beban kerja')
p('Gunakan laporan cetak, CSV, atau XLSX sesuai kebutuhan. Periksa periode, unit, dan filter status sebelum mengunduh. Data ekspor mencakup status validasi, kapasitas efektif dalam jam, waktu aktual, utilisasi, dan status beban.')
p('Laporan baris anggota tidak dibatasi hanya pada Disetujui. Karena itu, jumlah seluruh baris ekspor dapat berbeda dari kartu ringkasan tim. Untuk rekonsiliasi, gunakan periode dan filter yang sama, lalu jumlahkan hanya baris Disetujui.')
h('Laporan progres mingguan dan bulanan')
steps(['Buka Lihat laporan progres atau halaman laporan progres. Pilih anggota sesuai hak akses dan periode.',
'Pilih Bulanan untuk seluruh bulan. Untuk Mingguan, tentukan tanggal awal di dalam periode; rentang adalah tanggal awal sampai enam hari berikutnya, dibatasi akhir bulan.',
'Tinjau pekerjaan menurut kategori, status, ringkasan, kendala, langkah, dan target. Total durasi berasal dari aktivitas dalam rentang laporan.'])
p('Status pekerjaan mengikuti pelengkap progres terbaru dalam rentang yang dipilih. Bila aktivitas tidak memiliki pelengkap progres pada rentang itu, laporan menampilkannya sebagai Selesai dikerjakan dengan 100%. Lengkapi progres agar pekerjaan yang masih berjalan tidak terbaca selesai.')
p('Laporan progres dan utilisasi menjawab hal berbeda: progres menjelaskan hasil penyelesaian, sedangkan utilisasi membandingkan durasi kerja tercatat dengan kapasitas efektif. Progres 50% bukan berarti utilisasi 50%.')

page('7 Rumus hari kerja dan kapasitas efektif')
table(['Simbol', 'Arti dan satuan'], [
('K dan O', 'Jumlah hari kerja dan hari off dalam satu siklus'),
('A', 'Tanggal acuan sebagai hari kerja pertama siklus'),
('H', 'Jumlah hari kerja hasil perhitungan siklus dalam bulan'),
('J', 'Jam kerja per hari dari jadwal pengguna'),
('P', 'Persentase produktif periode, 1 sampai 100'),
('G dan E', 'Kapasitas kotor dan efektif, dalam menit'),
], [3, 13.8])
h('Menghitung hari kerja dari siklus')
p('Panjang siklus L = K + O. Untuk setiap tanggal d dari awal sampai akhir bulan, hitung selisih hari bertanda antara d dan A, kemudian posisinya dalam siklus:')
p('posisi(d) = ((selisih_hari(d, A) mod L) + L) mod L')
p('Tanggal d dihitung sebagai hari kerja jika posisi(d) < K. H adalah jumlah tanggal yang memenuhi kondisi tersebut. Posisi dimulai dari 0; normalisasi modulo membuat tanggal sebelum acuan tetap dihitung benar.')
p('Siklus berulang terus lintas bulan dan tidak dimulai ulang setiap tanggal 1. Akhir pekan tidak otomatis dianggap off. Hari libur nasional, cuti, dan lembur juga tidak dikoreksi otomatis oleh rumus siklus ini.')
h('Menghitung kapasitas')
p('G = round(H × J × 60)')
p('E = round(G × P / 100)')
p('Kapasitas efektif dalam jam = E / 60')
p('round berarti pembulatan ke menit bulat terdekat. Pembulatan G dilakukan sebelum menghitung E. Gunakan nilai menit tersimpan untuk perhitungan lanjutan agar tidak terjadi selisih dari pembulatan tampilan jam.')
h('Arti persentase produktif')
p('Jika P = 80%, kapasitas efektif adalah 80% dari kapasitas kotor. P merupakan parameter kebijakan periode, bukan angka hasil pengukuran progres Staff. Menurunkan P akan menaikkan utilisasi untuk total durasi yang sama.')
p('Acuan kapasitas otomatis dibuat dari jadwal pengguna dan standar periode. Staff tidak mengubahnya melalui formulir aktivitas. Pastikan K, O, A, J, dan P benar sebelum salinan kapasitas terbentuk.')

page('8 Rumus beban aktual dan klasifikasi')
h('Total waktu dan utilisasi individu')
p('Tᵢ = durasi aktual catatan ke i dalam menit')
p('B = jumlah seluruh Tᵢ pada pengajuan dalam periode')
p('U = round((B / E) × 100, 1), bila E > 0')
p('Beban dalam jam = B / 60. Total per sifat pekerjaan adalah penjumlahan Tᵢ untuk sifat yang sama. Rutin, Proyek, dan Mendadak tidak memiliki bobot pengali khusus; sifat tambahan juga dikelompokkan menurut nilai yang tersimpan.')
p('Jika E = 0 atau kapasitas belum ada, U tidak tersedia dan statusnya Kapasitas belum tersedia. Jika E > 0 dan belum ada aktivitas, U = 0,0%; periksa kelengkapan input sebelum menafsirkan hasil tersebut.')
h('Volume dan waktu rata rata')
p('Rata-rata menit per unit = round(Tᵢ / volume aktual, 2). Nilai ini merupakan informasi turunan. Beban yang dijumlahkan tetap Tᵢ, bukan hasil perkalian ulang volume dengan rata-rata yang sudah dibulatkan.')
p('Fungsi estimasi volume × rata-rata waktu masih ada dalam kode, tetapi alur pencatatan aktual menggunakan durasi yang diinput langsung. Jangan mencampur rumus estimasi dengan beban aktual dalam satu rekonsiliasi.')
h('Batas status utilisasi')
p('Misalkan a adalah batas Kapasitas tersedia, b batas atas Sehat, dan c batas atas Padat. Klasifikasi memakai U yang sudah dibulatkan satu desimal.')
table(['Status', 'Aturan', 'Nilai cadangan aplikasi'], [
('Kapasitas tersedia', 'U < a', 'U < 70%'),
('Sehat', 'a ≤ U ≤ b', '70% ≤ U ≤ 85%'),
('Padat', 'b < U ≤ c', '85% < U ≤ 100%'),
('Kelebihan beban', 'U > c', 'U > 100%'),
('Kapasitas belum tersedia', 'E tidak positif', 'Utilisasi tidak tersedia'),
], [5.8, 4.4, 6.6])
p('Sistem memilih ambang dengan tanggal berlaku terbaru yang tidak melewati tanggal awal periode. Jika tidak ada yang cocok, digunakan 70, 85, dan 100 dengan penanda sementara. Angka ini bukan standar regulasi atau penilaian kinerja universal.')

page('9 Contoh hitung individu dan tim')
h('Contoh individu pada September 2026')
p('Asumsi: siklus 5 hari kerja dan 2 hari off, acuan Senin 31 Agustus 2026, 8 jam per hari, produktif 80%, dan ambang cadangan 70/85/100. Rentang perhitungan 1–30 September 2026 menghasilkan 22 hari kerja.')
p('G = round(22 × 8 × 60) = 10.560 menit = 176 jam')
p('E = round(10.560 × 80 / 100) = 8.448 menit = 140,8 jam')
table(['Kelompok aktivitas', 'Total menit aktual', 'Jam'], [
('Rutin', '4.200', '70'), ('Proyek', '2.400', '40'), ('Mendadak', '600', '10'), ('Total B', '7.200', '120'),
], [7.8, 5, 4])
p('U = round((7.200 / 8.448) × 100, 1) = 85,2%')
p('Hasilnya Padat karena 85,2% lebih dari 85% dan tidak melebihi 100%. Sebagai analisis tambahan, selisih kapasitas adalah E − B = 1.248 menit atau 20,8 jam. Selisih ini bukan jaminan waktu bebas sebelum kelengkapan catatan diperiksa.')
h('Contoh ringkasan tim')
p('Rumus tim: U tim = round((jumlah B disetujui / jumlah E disetujui) × 100, 1). Gunakan kelompok anggota yang lolos filter dan berada dalam cakupan akses.')
table(['Anggota', 'Status', 'E menit', 'B menit', 'U'], [
('A', 'Disetujui', '8.448', '7.200', '85,2%'),
('B', 'Disetujui', '4.224', '2.400', '56,8%'),
('C', 'Diajukan', '8.448', '9.000', '106,5%'),
], [2.3, 4, 3.5, 3.5, 3.5])
p('U tim = round((9.600 / 12.672) × 100, 1) = 75,8%. Anggota C belum dihitung dalam ringkasan tim karena belum disetujui. Rata-rata sederhana (85,2% + 56,8%) / 2 = 71,0% tidak digunakan karena kapasitas A dan B berbeda.')
p('Batas penting: 70,0% adalah Sehat; 85,0% masih Sehat; 85,1% adalah Padat; 100,0% masih Padat; 100,1% adalah Kelebihan beban, jika memakai ambang cadangan.')

page('10 Pemeriksaan hasil dan penanganan kendala')
table(['Kondisi', 'Pemeriksaan dan tindakan'], [
('Belum ada periode terbuka', 'Minta PIC memeriksa status periode. Input membutuhkan periode open.'),
('Tombol pengajuan belum tersedia', 'Pengajuan pertama menunggu akhir bulan. Pastikan periode terbuka, kapasitas ada, dan minimal satu aktivitas tersedia.'),
('Data tidak bisa diubah', 'Cek status pengajuan dan periode. Diajukan/Disetujui perlu dikembalikan ke Perlu Revisi; periode terkunci perlu dibuka kembali.'),
('Tanggal ditolak', 'Gunakan tanggal aktivitas atau laporan dalam bulan aktif dan tidak setelah hari ini.'),
('Aktivitas ganda ditolak', 'Tinjau daftar sebelum mencatat ulang. Jangan membedakan nama hanya untuk menggandakan durasi.'),
('Utilisasi kosong', 'Periksa keberadaan kapasitas serta hari kerja hasil siklus. Nol kapasitas berbeda dari nol aktivitas.'),
('Ringkasan tim berbeda dari ekspor', 'Samakan periode dan filter, kemudian bandingkan hanya data Disetujui.'),
('Perubahan jadwal tidak mengubah angka', 'Kapasitas lama merupakan salinan tersimpan. Minta pengelola meninjau koreksi, jangan hapus data historis sembarangan.'),
('Pekerjaan tampak selesai padahal berjalan', 'Tambahkan pelengkap progres Sedang dikerjakan pada rentang laporan yang ditampilkan.'),
], [5.3, 11.5])
h('Daftar pemeriksaan sebelum menyetujui')
p('Periksa periode dan jadwal; hitung ulang G dan E; jumlahkan menit aktivitas; hitung U satu desimal; cocokkan tanggal berlaku ambang; periksa catatan ganda dan input terlambat; pastikan status progres sesuai; baru tentukan keputusan validasi.')
h('Pengamanan hasil laporan')
p('Simpan ekspor hanya pada lokasi yang diizinkan organisasi. Batasi penyebaran data individu. Gunakan menu Audit untuk menelusuri tindakan yang dicatat sistem. Untuk kendala, sampaikan periode, menu, pesan error, serta contoh perhitungan kepada pengelola tanpa mengirim kata sandi.')

page('Lampiran Rujukan implementasi')
p('Rujukan berikut membantu pengelola menyesuaikan manual ketika aplikasi berubah. Semua lokasi bersifat relatif terhadap akar proyek lavel-vwork.')
table(['Rujukan', 'Fungsi yang menjadi acuan'], [
('app/Services/WorkScheduleCalculator.php', 'Siklus dan jumlah hari kerja'),
('app/Services/WorkloadCalculator.php', 'Kapasitas, penjumlahan beban, utilisasi, klasifikasi'),
('app/Http/Controllers/WorkloadEntryController.php', 'Durasi aktual, volume, dan pembuatan kapasitas'),
('app/Models/UtilizationThreshold.php', 'Pemilihan ambang dan nilai cadangan'),
('app/Models/WorkloadSubmission.php', 'Syarat edit dan pengajuan'),
('app/Services/SubmissionWorkflow.php', 'Transisi validasi dan riwayat'),
('app/Services/VisibleTeam.php', 'Cakupan anggota per peran'),
('app/Http/Controllers/DashboardController.php', 'Agregasi tim berdasarkan data Disetujui'),
('app/Http/Controllers/ProgressReportController.php', 'Rentang dan pengelompokan laporan progres'),
('routes/web.php', 'Menu, endpoint, dan pembatasan peran'),
], [10, 6.8])
h('Istilah data untuk rekonsiliasi')
p('required_minutes menyimpan durasi aktual pada alur input saat ini. monthly_volume menyimpan volume aktual satu catatan. average_minutes_per_unit menyimpan hasil bagi durasi dan volume. Nama kolom historis tersebut tidak berarti pengguna harus mengisi estimasi bulanan.')
p('effective_minutes adalah penyebut utilisasi. Status pengajuan (Draft, Diajukan, Perlu Revisi, Disetujui) berbeda dari status beban (Kapasitas tersedia, Sehat, Padat, Kelebihan beban) dan status progres (Selesai, Sedang, Akan dikerjakan).')
h('Pemeliharaan manual')
p('Tinjau ulang manual bila aturan kapasitas, peran, form input, ambang, atau alur validasi berubah. Simpan versi sebelumnya agar pembaca laporan historis dapat mengenali aturan yang dipakai pada masanya.')

# Verify the published example independently of document formatting.
anchor = date(2026, 8, 31)
days = sum(((date(2026, 9, 1) + timedelta(days=i) - anchor).days % 7) < 5 for i in range(30))
round1 = lambda x: Decimal(str(x)).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)
assert days == 22
assert days * 8 * 60 == 10560
assert Decimal(10560) * Decimal('0.8') == 8448
assert round1(Decimal(7200) / Decimal(8448) * 100) == Decimal('85.2')
assert round1(Decimal(9600) / Decimal(12672) * 100) == Decimal('75.8')
assert round1(Decimal(9000) / Decimal(8448) * 100) == Decimal('106.5')
doc.save(OUT)
print(OUT)
print(f'Example checks passed; paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}')
